import unittest
from datetime import datetime
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile
from fastapi.testclient import TestClient
from openpyxl import Workbook
from main import app
from workbook_import import MAX_UPLOAD_BYTES, WorkbookError, parse_workbook


def workbook_bytes(edit=None):
    w = Workbook()
    w.remove(w.active)
    for year, offset, h in ((2026, 2, 3), (2027, 4, 2)):
        s = w.create_sheet(f"Jul {year} - Jun {year + 1}")
        for col, label in enumerate(("Name", "MCR No.", "Residency Year"), offset):
            s.cell(h, col, label)
        for i in range(12):
            col = offset + 3 + i
            s.cell(h - 1, col, datetime(year + (i + 6) // 12, (i + 6) % 12 + 1, 1))
            s.cell(h, col, "7 Jul - 5 Aug (original display text)")
            s.cell(h + 1, col, "LOA" if i == 0 else "Posting")
        for col, value in enumerate(("Example Resident", "M12345A", "R2,R3"), offset):
            s.cell(h + 1, col, value)
        if year == 2027:
            s.cell(h - 1, offset + 15, "Remarks")
            s.merge_cells(start_row=h - 1, end_row=h, start_column=offset + 15, end_column=offset + 15)
            s.cell(h + 1, offset + 15, "Preserved merged header remark")
            s.merge_cells(start_row=h + 1, end_row=h + 1, start_column=offset + 4, end_column=offset + 5)
    s = w.create_sheet("Request")
    s.append(["Name", "MCR No.", "Residency Year", "Posting Request", "Period", "Remarks"])
    s.append(["Example Resident", "M12345A", "R2,R3", "Free text request", "Either Jul or Oct", "Unconfirmed"])
    s = w.create_sheet("Postings")
    s.append(["R2 Electives", None, None, "Residency Year", "Core Posting", "Duration (Months)", "Max", "Remarks"])
    s.append(["Raw elective note", None, None, "R2", "Core example", 3, 2, "Original condition"])
    if edit:
        edit(w)
    s.append([None, None, None, None, "Standalone posting", 6])
    stream = BytesIO()
    w.save(stream)
    return stream.getvalue()


class WorkbookImportTests(unittest.TestCase):
    def test_history_matches_mcr_preserves_blanks_and_excludes_future(self):
        def edit(w):
            target = w["Jul 2027 - Jun 2028"]
            target["D3"] = "  EXAMPLE   Resident "
            target["E3"] = " m12345a "
            target["R3"] = None
            future = w.copy_worksheet(target)
            future.title = "Jul 2028 - Jun 2029"
            for cell in future[1]:
                if isinstance(cell.value, datetime):
                    cell.value = cell.value.replace(year=cell.value.year + 1)
        r = parse_workbook(workbook_bytes(edit))
        self.assertEqual([h["year"] for h in r["histories"]], [2026])
        history = r["histories"][0]["residents"][0]
        self.assertEqual(history["match_status"], "matched")
        self.assertEqual(history["matched_resident_source"], r["residents"][0]["source"])
        self.assertEqual(r["requests"][0]["match_status"], "matched")
        self.assertEqual(r["residents"][0]["mcr"], " m12345a ")
        self.assertEqual(len(r["residents"][0]["postings"]), 12)
        self.assertIsNone(r["residents"][0]["postings"][-1]["value"])
        self.assertEqual(r["months"][1]["month"], "2027-08")
        self.assertEqual(r["months"][1]["date_range_text"], "7 Jul - 5 Aug (original display text)")
        unmapped = next(i for i in r["issues"] if i["code"] == "unmapped_posting")
        self.assertEqual(unmapped["severity"], "info")
        self.assertIn("LOA", unmapped["values"])

    def test_identity_problems_withhold_matching_without_dropping_rows(self):
        def duplicate(w):
            s = w["Jul 2027 - Jun 2028"]
            s["D4"], s["E4"], s["F4"] = "Example Resident", "M12345A", "R3"
        cases = [
            (lambda w: setattr(w["Jul 2027 - Jun 2028"]["E3"], "value", None), "missing_mcr"),
            (duplicate, "duplicate_mcr"),
            (lambda w: setattr(w["Jul 2026 - Jun 2027"]["B4"], "value", "Different Name"), "conflicting_identity"),
            (lambda w: setattr(w["Jul 2026 - Jun 2027"]["C4"], "value", "MP1234A"), "conflicting_identity"),
            (lambda w: setattr(w["Jul 2026 - Jun 2027"]["C4"], "value", None), "missing_mcr"),
        ]
        for edit, code in cases:
            with self.subTest(code=code, edit=edit):
                r = parse_workbook(workbook_bytes(edit))
                history = r["histories"][0]["residents"][0]
                self.assertNotEqual(history["match_status"], "matched")
                self.assertIsNone(history["matched_resident_source"])
                self.assertIn(code, [i["code"] for i in r["issues"]])
                self.assertEqual(len(history["postings"]), 12)

    def test_months_history_and_raw_text_remain_separate(self):
        r = parse_workbook(workbook_bytes())
        self.assertEqual([s["year"] for s in r["available_years"]], [2026, 2027])
        self.assertEqual(r["months"][0]["month"], "2027-07")
        self.assertEqual(r["months"][-1]["month"], "2028-06")
        self.assertEqual(r["months"][0]["date_range_text"], "7 Jul - 5 Aug (original display text)")
        self.assertEqual(r["residents"][0]["residency_year"], "R2,R3")
        p = r["residents"][0]["postings"]
        self.assertEqual(r["residents"][0]["remarks"], "Preserved merged header remark")
        self.assertEqual(p[0]["value"], "LOA")
        self.assertEqual(p[1]["source"], p[2]["source"])
        self.assertEqual(r["histories"][0]["year"], 2026)
        self.assertNotIn("completed", r["histories"][0]["residents"][0])
        self.assertEqual(r["requests"][0]["period"], "Either Jul or Oct")
        self.assertEqual(r["postings"]["core"][0]["duration_months"], "3")
        self.assertEqual(r["postings"]["rows"][1]["cells"]["A"], "Raw elective note")
        self.assertEqual(r["postings"]["core"][1]["posting"], "Standalone posting")

    def test_missing_target_and_corrupt_workbooks(self):
        with self.assertRaisesRegex(WorkbookError, "No July"):
            parse_workbook(workbook_bytes(), 2030)
        with self.assertRaisesRegex(WorkbookError, "valid, unencrypted"):
            parse_workbook(b"not a workbook")
        stream = BytesIO()
        with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
            archive.writestr("large.xml", b"x" * (51 * 1024 * 1024))
        with self.assertRaisesRegex(WorkbookError, "expanded"):
            parse_workbook(stream.getvalue())

    def test_api_upload_validation(self):
        with TestClient(app) as client:
            good = client.post("/api/import", files={"file": ("example.xlsx", workbook_bytes())})
            self.assertEqual(good.status_code, 200, good.text)
            self.assertEqual(good.json()["target_year"], 2027)
            for name, content in (("wrong.csv", b"text"), ("broken.xlsx", b"bad")):
                self.assertEqual(client.post("/api/import", files={"file": (name, content)}).status_code, 400)
            self.assertEqual(client.post("/api/import", files={"file": ("large.xlsx", b"x" * (MAX_UPLOAD_BYTES + 1))}).status_code, 413)
            self.assertEqual(client.post("/api/import", data={"target_year": "bad"}, files={"file": ("example.xlsx", workbook_bytes())}).status_code, 422)


if __name__ == "__main__":
    unittest.main()
