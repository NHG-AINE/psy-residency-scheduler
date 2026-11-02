"""Read posting records without interpreting them as completed training or rules."""
from datetime import date, datetime
from io import BytesIO
from zipfile import ZipFile

from openpyxl import load_workbook

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 50 * 1024 * 1024


class WorkbookError(ValueError):
    pass


def text(value):
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def source(sheet, row, column=None):
    ref = {"sheet": sheet.title, "row": row}
    if column is not None:
        ref["column"] = column
    return ref


def header(sheet, required):
    for row in sheet.iter_rows(max_row=min(sheet.max_row, 10)):
        columns = {str(c.value).strip().lower(): c.column for c in row if c.value is not None}
        if required <= columns.keys():
            for merged in sheet.merged_cells.ranges:
                if merged.min_row <= row[0].row <= merged.max_row:
                    value = sheet.cell(merged.min_row, merged.min_col).value
                    if value is not None:
                        columns.setdefault(str(value).strip().lower(), merged.min_col)
            return row[0].row, columns
    raise WorkbookError(f"{sheet.title}: required headers missing: {', '.join(sorted(required))}")


def records(sheet, header_row, columns, fields):
    result = []
    for row in range(header_row + 1, sheet.max_row + 1):
        record = {field: text(sheet.cell(row, columns[label]).value) if label in columns else None
                  for field, label in fields.items()}
        if any(record.get(key) for key in ("name", "mcr", "residency_year", "posting")):
            record["source"] = source(sheet, row)
            result.append(record)
    return result


IDENTITY_FIELDS = {"name": "name", "mcr": "mcr no.", "residency_year": "residency year",
                   "intake": "residency intake", "mom_number": "mom number", "email": "email", "remarks": "remarks"}


def yearly_sheet(sheet):
    header_row, columns = header(sheet, {"name", "mcr no.", "residency year"})
    month_cells = []
    for row in sheet.iter_rows(max_row=header_row):
        dates = [c for c in row if isinstance(c.value, (date, datetime))]
        if len(dates) == 12:
            month_cells = dates
            break
    if not month_cells:
        raise WorkbookError(f"{sheet.title}: twelve calendar month headings are required")
    first = month_cells[0].value
    expected = [f"{first.year + (i + 6) // 12:04d}-{(i + 6) % 12 + 1:02d}" for i in range(12)]
    if first.month != 7 or [c.value.strftime("%Y-%m") for c in month_cells] != expected:
        raise WorkbookError(f"{sheet.title}: headings must run July through June in order")
    months = [{"month": c.value.strftime("%Y-%m"),
               "date_range_text": text(sheet.cell(header_row, c.column).value),
               "source": source(sheet, c.row, c.column_letter)} for c in month_cells]
    # The target sheet has merged SAF notes spanning the year. Retain the anchor
    # reference and text across its months; do not interpret the note as training.
    anchors = {}
    for merged in sheet.merged_cells.ranges:
        for row in range(merged.min_row, merged.max_row + 1):
            for col in range(merged.min_col, merged.max_col + 1):
                anchors[row, col] = (merged.min_row, merged.min_col)
    residents = records(sheet, header_row, columns, IDENTITY_FIELDS)
    for resident in residents:
        row = resident["source"]["row"]
        resident["postings"] = []
        for c, month in zip(month_cells, months):
            cell = sheet.cell(*anchors.get((row, c.column), (row, c.column)))
            resident["postings"].append({"month": month["month"], "value": text(cell.value),
                                          "source": source(sheet, cell.row, cell.column_letter)})
    return {"year": first.year, "sheet": sheet.title, "months": months, "residents": residents}


def parse_workbook(content: bytes, target_year: int = 2027) -> dict:
    if len(content) > MAX_UPLOAD_BYTES:
        raise WorkbookError("Workbook exceeds the 10 MB upload limit")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_EXPANDED_BYTES:
                raise WorkbookError("Workbook expanded contents exceed 50 MB")
            if len(archive.infolist()) > 2000:
                raise WorkbookError("Workbook contains too many archive entries")
        workbook = load_workbook(BytesIO(content), data_only=True, keep_links=False)
    except WorkbookError:
        raise
    except Exception as exc:
        raise WorkbookError("Unable to read this XLSX workbook; upload a valid, unencrypted .xlsx file") from exc
    try:
        for sheet in workbook:
            if sheet.max_row > 10000 or sheet.max_column > 100:
                raise WorkbookError(f"{sheet.title}: worksheet exceeds supported dimensions (10000 rows, 100 columns)")
        years = sorted([yearly_sheet(s) for s in workbook if s.title.lower().startswith("jul")], key=lambda s: s["year"])
        if len({s["year"] for s in years}) != len(years):
            raise WorkbookError("Workbook contains ambiguous duplicate planning years")
        target = next((s for s in years if s["year"] == target_year), None)
        if target is None:
            raise WorkbookError(f"No July {target_year} - June {target_year + 1} worksheet found")
        for required in ("Request", "Postings"):
            if required not in workbook.sheetnames:
                raise WorkbookError(f"Required worksheet missing: {required}")
        request = workbook["Request"]
        request_row, request_columns = header(request, {"name", "mcr no.", "posting request", "period"})
        requests = records(request, request_row, request_columns,
                           {**IDENTITY_FIELDS, "posting_request": "posting request", "period": "period"})
        postings = workbook["Postings"]
        posting_row, posting_columns = header(postings, {"residency year", "core posting", "duration (months)", "max", "remarks"})
        core = records(postings, posting_row, posting_columns,
                       {"residency_year": "residency year", "posting": "core posting", "duration_months": "duration (months)", "max": "max", "remarks": "remarks"})
        rows = [{"source": source(postings, row[0].row),
                 "cells": {c.column_letter: text(c.value) for c in row if c.value is not None}}
                for row in postings if any(c.value is not None for c in row)]
        return {"target_year": target_year, "target_sheet": target["sheet"],
                "available_years": [{"year": s["year"], "sheet": s["sheet"]} for s in years],
                "months": target["months"], "residents": target["residents"],
                "histories": [s for s in years if s["year"] < target_year], "requests": requests,
                "postings": {"core": core, "rows": rows}, "warnings": []}
    finally:
        workbook.close()

